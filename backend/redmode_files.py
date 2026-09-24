"""Private RedMode file storage and safe scope-document extraction."""

from __future__ import annotations

import csv
import io
import json
import zipfile

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from bson import ObjectId
from gridfs.errors import NoFile
from motor.motor_asyncio import AsyncIOMotorGridFSBucket


TEXT_EXTENSIONS = {".txt", ".csv", ".json"}
DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".xlsx"}
ALLOWED_EXTENSIONS = TEXT_EXTENSIONS | DOCUMENT_EXTENSIONS
MAX_UNCOMPRESSED_DOCUMENT_BYTES = 25 * 1024 * 1024
MAX_EXTRACTED_CHARACTERS = 2_000_000


def clean_filename(filename: str) -> str:
    return filename.replace("\\", "/").split("/")[-1].replace("\r", "").replace("\n", "")[:200]


def extract_text(filename: str, content: bytes) -> str:
    name = clean_filename(filename)
    extension = "." + name.rsplit(".", 1)[-1].lower() if "." in name else ""
    if extension not in TEXT_EXTENSIONS:
        raise ValueError("scope_file_type_not_supported")
    try:
        decoded = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("scope_file_not_utf8") from exc
    if "\x00" in decoded:
        raise ValueError("scope_file_invalid_content")
    if extension == ".json":
        try:
            return json.dumps(json.loads(decoded), ensure_ascii=False, indent=2)
        except json.JSONDecodeError as exc:
            raise ValueError("scope_file_invalid_json") from exc
    if extension == ".csv":
        try:
            rows = csv.reader(io.StringIO(decoded), strict=True)
            return "\n".join(" ".join(cell for cell in row) for row in rows)
        except csv.Error as exc:
            raise ValueError("scope_file_invalid_csv") from exc
    return decoded


def _check_zip_size(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if len(archive.infolist()) > 5_000 or sum(item.file_size for item in archive.infolist()) > MAX_UNCOMPRESSED_DOCUMENT_BYTES:
                raise ValueError("scope_document_too_large_uncompressed")
    except zipfile.BadZipFile as exc:
        raise ValueError("scope_file_unreadable") from exc


def _flatten_sections(sections: list[tuple[str, str]]) -> tuple[str, dict[int, str]]:
    lines = []
    positions = {}
    char_count = 0
    for position, section in sections:
        for line in section.splitlines():
            if not line.strip():
                continue
            char_count += len(line)
            if char_count > MAX_EXTRACTED_CHARACTERS:
                raise ValueError("scope_document_too_much_text")
            lines.append(line)
            positions[len(lines)] = position
    if not lines:
        raise ValueError("scope_file_no_text")
    return "\n".join(lines), positions


def extract_scope_content(filename: str, content: bytes) -> tuple[str, dict[int, str]]:
    """Return extracted text and the available source position of each line."""
    name = clean_filename(filename)
    extension = "." + name.rsplit(".", 1)[-1].lower() if "." in name else ""
    if extension in TEXT_EXTENSIONS:
        return extract_text(name, content), {}
    if extension not in DOCUMENT_EXTENSIONS:
        raise ValueError("scope_file_type_not_supported")

    try:
        if extension == ".pdf":
            reader = PdfReader(io.BytesIO(content), strict=False)
            if reader.is_encrypted:
                raise ValueError("scope_file_protected")
            if len(reader.pages) > 200:
                raise ValueError("scope_document_too_many_pages")
            return _flatten_sections([
                (f"página {number}", page.extract_text() or "")
                for number, page in enumerate(reader.pages, start=1)
            ])

        _check_zip_size(content)
        if extension == ".docx":
            document = Document(io.BytesIO(content))
            sections = [(f"parágrafo {number}", paragraph.text) for number, paragraph in enumerate(document.paragraphs, start=1)]
            for table_number, table in enumerate(document.tables, start=1):
                for row_number, row in enumerate(table.rows, start=1):
                    sections.append((f"tabela {table_number}, linha {row_number}", " ".join(cell.text for cell in row.cells)))
            return _flatten_sections(sections)

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            sections = []
            if len(workbook.worksheets) > 100:
                raise ValueError("scope_document_too_many_sheets")
            for sheet in workbook.worksheets:
                if (sheet.max_row or 0) * (sheet.max_column or 0) > 50_000:
                    raise ValueError("scope_document_too_many_cells")
                for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                    if row_number > 50_000 or len(row) > 50_000 or len(sections) > 50_000:
                        raise ValueError("scope_document_too_many_cells")
                    value = " ".join(str(cell) for cell in row if cell is not None)
                    if value:
                        sections.append((f"planilha {sheet.title}, linha {row_number}", value))
            return _flatten_sections(sections)
        finally:
            workbook.close()
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("scope_file_unreadable") from exc


class GridFSPrivateStore:
    def __init__(self, db, bucket_name: str):
        self.bucket = AsyncIOMotorGridFSBucket(db, bucket_name=bucket_name)

    async def save(self, filename: str, content: bytes, metadata: dict) -> str:
        file_id = await self.bucket.upload_from_stream(filename, content, metadata=metadata)
        return str(file_id)

    async def read(self, file_id: str) -> bytes:
        try:
            stream = await self.bucket.open_download_stream(ObjectId(file_id))
            return await stream.read()
        except (NoFile, ValueError) as exc:
            raise FileNotFoundError(file_id) from exc

    async def delete(self, file_id: str) -> None:
        try:
            await self.bucket.delete(ObjectId(file_id))
        except NoFile:
            pass


class GridFSScopeStore(GridFSPrivateStore):
    def __init__(self, db):
        super().__init__(db, "redmode_scope_files")


class GridFSEvidenceStore(GridFSPrivateStore):
    def __init__(self, db):
        super().__init__(db, "redmode_evidence_files")
