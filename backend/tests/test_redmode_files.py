"""Scope-file parsing and source precedence."""

import io

import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from redmode_files import clean_filename, extract_scope_content, extract_text
from redmode_scope import compile_scope_text, merge_scope_rules


def test_clean_filename_removes_path_components():
    assert clean_filename("../../secret/scope.txt") == "scope.txt"
    assert clean_filename("C:\\secret\\scope.csv") == "scope.csv"


def test_json_categories_are_preserved_in_extracted_text():
    extracted = extract_text(
        "scope.json",
        b'{"client":["192.0.2.10"],"third_party":["203.0.113.10"],"excluded":["198.51.100.0/24"]}',
    )
    rules = compile_scope_text(extracted, source_id="file-1")
    assert {rule["value"]: rule["category"] for rule in rules} == {
        "192.0.2.10": "client",
        "203.0.113.10": "third_party",
        "198.51.100.0/24": "excluded",
    }


def test_csv_categories_and_merge_keep_source_provenance():
    csv_text = extract_text(
        "scope.csv",
        b"category,target\nthird_party,203.0.113.10\nclient,192.0.2.10\nexcluded,192.0.2.10",
    )
    rules = merge_scope_rules([
        compile_scope_text("192.0.2.10", source_id="text"),
        compile_scope_text(csv_text, source_id="file-1"),
    ])
    target = next(rule for rule in rules if rule["value"] == "192.0.2.10")
    assert target["category"] == "excluded"
    assert {origin["source_id"] for origin in target["origins"]} == {"text", "file-1"}


@pytest.mark.parametrize("filename,content", [
    ("scope.pdf", b"192.0.2.10"),
    ("scope.txt", b"\xff\xfe"),
    ("scope.json", b"{"),
    ("scope.csv", b"\x00"),
])
def test_invalid_files_are_rejected(filename, content):
    with pytest.raises(ValueError):
        extract_text(filename, content)


def test_docx_extracts_paragraph_and_table_positions():
    document = Document()
    document.add_paragraph("192.0.2.10")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "203.0.113.10"
    output = io.BytesIO()
    document.save(output)
    extracted, positions = extract_scope_content("scope.docx", output.getvalue())
    rules = compile_scope_text(extracted)
    assert {rule["value"] for rule in rules} == {"192.0.2.10", "203.0.113.10"}
    assert "parágrafo" in positions[1]
    assert "tabela" in positions[2]


def test_xlsx_extracts_sheet_and_row():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Escopo"
    sheet["A1"] = "https://app.example.test"
    output = io.BytesIO()
    workbook.save(output)
    extracted, positions = extract_scope_content("scope.xlsx", output.getvalue())
    assert compile_scope_text(extracted)[0]["kind"] == "url"
    assert positions[1] == "planilha Escopo, linha 1"


def test_pdf_extracts_page_position():
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)}),
    })
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 72 720 Td (192.0.2.10) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    extracted, positions = extract_scope_content("scope.pdf", output.getvalue())
    assert compile_scope_text(extracted)[0]["value"] == "192.0.2.10"
    assert positions[1] == "página 1"


def test_image_only_pdf_has_clear_error():
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    output = io.BytesIO()
    writer.write(output)
    with pytest.raises(ValueError, match="scope_file_no_text"):
        extract_scope_content("scan.pdf", output.getvalue())
