"""Opt-in RedMode HTTP and GridFS checks against an isolated local MongoDB."""

import io
import os
from uuid import uuid4

import pytest
from docx import Document
from httpx import ASGITransport, AsyncClient
from motor.motor_asyncio import AsyncIOMotorClient
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

import app_state
from auth import create_access_token
from db import db_manager
from main import app


def _headers(username):
    token = create_access_token({"sub": username, "role": "tech"})
    return {"Authorization": f"Bearer {token}"}


def _scope_uploads():
    document = Document()
    document.add_paragraph("192.0.2.22")
    docx = io.BytesIO()
    document.save(docx)

    workbook = Workbook()
    workbook.active["A1"] = "192.0.2.23"
    xlsx = io.BytesIO()
    workbook.save(xlsx)

    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)}),
    })
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 72 720 Td (192.0.2.21) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    pdf = io.BytesIO()
    writer.write(pdf)

    return [
        ("scope.txt", b"192.0.2.20", "text/plain"),
        ("scope.csv", b"category,target\nthird_party,203.0.113.20", "text/csv"),
        ("scope.json", b'{"excluded":["198.51.100.20"]}', "application/json"),
        ("scope.pdf", pdf.getvalue(), "application/pdf"),
        ("scope.docx", docx.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        ("scope.xlsx", xlsx.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    ]


@pytest.mark.asyncio
async def test_redmode_persists_private_project_data_in_mongo(monkeypatch):
    uri = os.environ.get("REDMODE_TEST_MONGO_URI")
    if not uri:
        pytest.skip("Set REDMODE_TEST_MONGO_URI for an isolated local MongoDB")
    if uri != "mongodb://127.0.0.1:27019":
        pytest.fail("Integration test accepts only mongodb://127.0.0.1:27019")

    client = AsyncIOMotorClient(uri, serverSelectionTimeoutMS=5000, tz_aware=True)
    db_name = f"redmode_h0_test_{uuid4().hex}"
    db = client[db_name]
    await client.admin.command("ping")
    monkeypatch.setattr(db_manager, "db", db)
    monkeypatch.setattr(app_state, "APP_INITIALIZED", True)

    try:
        await db.users.insert_many([
            {"username": username, "role": "tech", "is_active": True,
             "extra_permissions": ["redmode:access"]}
            for username in ("h0-owner", "h0-member", "h0-outsider")
        ])
        owner = _headers("h0-owner")
        member = _headers("h0-member")
        outsider = _headers("h0-outsider")
        base = "/api/redmode/projects/h0-synthetic"

        # ASGITransport does not run app lifespan; no scheduler or scan worker starts.
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as http:
            created = await http.post("/api/redmode/projects", json={
                "slug": "h0-synthetic", "display_name": "H0 Synthetic",
            }, headers=owner)
            assert created.status_code == 201, created.text
            assert created.json()["members"] == ["h0-owner"]
            assert (await http.get(base, headers=outsider)).status_code == 403
            feed = (await http.get("/api/redmode/projects", headers=outsider)).json()
            assert set(feed["items"][0]) == {
                "slug", "display_name", "phase", "status", "responsible", "last_activity_at",
            }

            added = await http.put(f"{base}/members/h0-member", headers=owner)
            assert added.status_code == 200, added.text

            text_scope = await http.post(f"{base}/scope/text", json={
                "text": "192.0.2.10\nterceiro: 203.0.113.10\nexcluído: 198.51.100.0/24",
            }, headers=member)
            assert text_scope.status_code == 201, text_scope.text
            first_version = text_scope.json()["id"]

            uploads = _scope_uploads()
            bundle = await http.post(f"{base}/scope/submit", data={"text": "https://app.example.test"},
                                     files=[("files", item) for item in uploads], headers=member)
            assert bundle.status_code == 201, bundle.text
            active = bundle.json()
            assert active["id"] != first_version
            assert {item["value"] for item in active["rules"]} == {
                "https://app.example.test", "192.0.2.20", "203.0.113.20",
                "198.51.100.20", "192.0.2.21", "192.0.2.22", "192.0.2.23",
            }
            categories = {item["value"]: item["category"] for item in active["rules"]}
            assert categories["203.0.113.20"] == "third_party"
            assert categories["198.51.100.20"] == "excluded"
            assert len(active["source"]["files"]) == len(uploads)
            sources_response = await http.get(
                f"{base}/scope/versions/{active['id']}/sources",
                headers=member,
            )
            assert sources_response.status_code == 200, sources_response.text
            sources = sources_response.json()["items"]
            assert len(sources) == len(uploads) + 1
            assert sources[0]["source_id"] == "text"
            assert all(item["representation"]["materialized"] for item in sources)
            extracted = await http.get(
                f"{base}/scope/versions/{active['id']}/sources/"
                f"{sources[1]['source_id']}/representation",
                headers=member,
            )
            assert extracted.status_code == 200
            assert extracted.content == b"192.0.2.20"
            inventory = await http.get(
                f"{base}/scope/versions/{active['id']}/assets?limit=2",
                headers=member,
            )
            assert inventory.status_code == 200, inventory.text
            assert inventory.json()["total"] == 7
            assert len(inventory.json()["items"]) == 2
            file_id = active["source"]["files"][0]["id"]
            scope_file_url = f"{base}/scope/versions/{active['id']}/files/{file_id}"
            for metadata, (filename, content, _) in zip(active["source"]["files"], uploads):
                downloaded = await http.get(
                    f"{base}/scope/versions/{active['id']}/files/{metadata['id']}", headers=member,
                )
                assert metadata["filename"] == filename
                assert downloaded.status_code == 200 and downloaded.content == content
                assert downloaded.headers["cache-control"] == "private, no-store"
            assert (await http.get(scope_file_url, headers=outsider)).status_code == 403

            bad = await http.post(f"{base}/scope/submit", files=[
                ("files", ("bad.json", b"{", "application/json")),
            ], headers=member)
            assert bad.status_code == 422
            assert (await http.get(f"{base}/scope/active", headers=owner)).json()["id"] == active["id"]

            evidence = await http.post(f"{base}/evidence", data={
                "phase": "reconnaissance", "text": "Observação sintética", "target": "192.0.2.20",
            }, files={"file": ("proof.txt", b"synthetic evidence", "text/plain")}, headers=member)
            assert evidence.status_code == 201, evidence.text
            evidence_id = evidence.json()["id"]
            evidence_file_url = f"{base}/evidence/{evidence_id}/file"
            downloaded = await http.get(evidence_file_url, headers=owner)
            assert downloaded.status_code == 200 and downloaded.content == b"synthetic evidence"
            assert (await http.get(evidence_file_url, headers=outsider)).status_code == 403

            payload = {"title": "Achado sintético", "description": "Teste de persistência",
                       "severity": "low", "phase": "reconnaissance", "targets": ["192.0.2.20"],
                       "evidence_ids": [evidence_id]}
            finding = await http.post(f"{base}/findings", json=payload, headers=member)
            assert finding.status_code == 201, finding.text
            finding_id = finding.json()["id"]
            previous_revision = finding.json()["revision"]["id"]
            edited = await http.put(f"{base}/findings/{finding_id}", json={
                **payload, "description": "Segunda revisão", "expected_revision_id": previous_revision,
            }, headers=owner)
            assert edited.status_code == 200, edited.text
            assert edited.json()["revision"]["number"] == 2
            revisions = (await http.get(f"{base}/findings/{finding_id}/revisions", headers=member)).json()
            assert [item["number"] for item in revisions["items"]] == [2, 1]
            linked = await http.post(f"{base}/evidence", data={
                "phase": "reporting", "text": "Evidência vinculada depois", "finding_id": finding_id,
            }, headers=owner)
            assert linked.status_code == 201, linked.text
            current_finding = (await http.get(f"{base}/findings/{finding_id}", headers=member)).json()
            assert linked.json()["id"] in current_finding["evidence_ids"]

            # Reopen the database client to detect accidental in-process-only state.
            client.close()
            client = AsyncIOMotorClient(uri, serverSelectionTimeoutMS=5000, tz_aware=True)
            db = client[db_name]
            monkeypatch.setattr(db_manager, "db", db)
            assert (await http.get(f"{base}/scope/active", headers=member)).json()["id"] == active["id"]
            assert (await http.get(scope_file_url, headers=owner)).content == b"192.0.2.20"
            assert (await http.get(evidence_file_url, headers=member)).content == b"synthetic evidence"
            assert (await http.get(f"{base}/findings/{finding_id}", headers=owner)).json()["revision"]["number"] == 2

            removed = await http.delete(f"{base}/members/h0-member", headers=owner)
            assert removed.status_code == 200, removed.text
            for url in (base, f"{base}/scope/active", scope_file_url, f"{base}/evidence",
                        evidence_file_url, f"{base}/findings", f"{base}/findings/{finding_id}"):
                assert (await http.get(url, headers=member)).status_code == 403, url
    finally:
        await client.drop_database(db_name)
        client.close()
