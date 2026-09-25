"""First RedMode slice: workspace permission, project feed, and creator access."""

import io
from datetime import datetime, timezone

import pytest
from docx import Document

from auth import create_access_token


class MemoryScopeFiles:
    def __init__(self):
        self.files = {}
        self.fail_save = False
        self.fail_after = None

    async def save(self, filename, content, metadata):
        if self.fail_save or (self.fail_after is not None and len(self.files) >= self.fail_after):
            raise RuntimeError("storage unavailable")
        file_id = f"file-{len(self.files) + 1}"
        self.files[file_id] = content
        return file_id

    async def read(self, file_id):
        if file_id not in self.files:
            raise FileNotFoundError(file_id)
        return self.files[file_id]

    async def delete(self, file_id):
        self.files.pop(file_id, None)


@pytest.fixture
def scope_files(monkeypatch):
    import routers.redmode as redmode
    store = MemoryScopeFiles()
    monkeypatch.setattr(redmode, "file_store", lambda: store)
    return store


@pytest.fixture(autouse=True)
def scope_source_texts(monkeypatch):
    import routers.redmode as redmode
    store = MemoryScopeFiles()
    monkeypatch.setattr(redmode, "source_text_store", lambda: store)
    return store


def headers_for(username: str, role: str = "tech") -> dict[str, str]:
    token = create_access_token({"sub": username, "role": role})
    return {"Authorization": f"Bearer {token}"}


async def grant_redmode(fake_db, username: str) -> None:
    await fake_db.users.update_one(
        {"username": username},
        {"$set": {"extra_permissions": ["redmode:access"]}},
    )


@pytest.mark.asyncio
async def test_workspace_requires_redmode_permission(async_client):
    response = await async_client.get("/api/redmode/projects", headers=headers_for("techuser"))
    assert response.status_code == 403

    response = await async_client.post(
        "/api/redmode/projects",
        json={"slug": "cliente-demo", "display_name": "Cliente Demo"},
        headers=headers_for("techuser"),
    )
    assert response.status_code == 403


@pytest.mark.asyncio
async def test_creator_is_member_and_feed_is_summary_only(async_client, fake_db):
    await grant_redmode(fake_db, "techuser")
    created = await async_client.post(
        "/api/redmode/projects",
        json={"slug": "cliente-demo", "display_name": "Cliente Demo"},
        headers=headers_for("techuser"),
    )
    assert created.status_code == 201
    body = created.json()
    assert body["slug"] == "cliente-demo"
    assert body["members"] == ["techuser"]
    assert body["responsible"] == "techuser"

    detail = await async_client.get(
        "/api/redmode/projects/cliente-demo", headers=headers_for("techuser")
    )
    assert detail.status_code == 200
    assert detail.json()["members"] == ["techuser"]

    feed = await async_client.get("/api/redmode/projects", headers=headers_for("techuser"))
    assert feed.status_code == 200
    assert feed.json()["total"] == 1
    item = feed.json()["items"][0]
    assert set(item) == {"slug", "display_name", "phase", "status", "responsible", "last_activity_at"}
    assert item["slug"] == "cliente-demo"
    assert "members" not in item


@pytest.mark.asyncio
async def test_enabled_nonmember_sees_summary_but_not_detail(async_client, fake_db):
    created = await async_client.post(
        "/api/redmode/projects",
        json={"slug": "cliente-demo", "display_name": "Cliente Demo"},
        headers=headers_for("admin", "admin"),
    )
    assert created.status_code == 201
    await grant_redmode(fake_db, "techuser")

    feed = await async_client.get("/api/redmode/projects", headers=headers_for("techuser"))
    assert feed.status_code == 200
    assert feed.json()["total"] == 1
    detail = await async_client.get(
        "/api/redmode/projects/cliente-demo", headers=headers_for("techuser")
    )
    assert detail.status_code == 403


@pytest.mark.asyncio
async def test_admin_without_membership_cannot_open_project(async_client, fake_db):
    await grant_redmode(fake_db, "techuser")
    created = await async_client.post(
        "/api/redmode/projects",
        json={"slug": "cliente-demo", "display_name": "Cliente Demo"},
        headers=headers_for("techuser"),
    )
    assert created.status_code == 201

    detail = await async_client.get(
        "/api/redmode/projects/cliente-demo", headers=headers_for("admin", "admin")
    )
    assert detail.status_code == 403


@pytest.mark.asyncio
async def test_project_feed_can_be_paginated(async_client):
    for slug in ("cliente-um", "cliente-dois"):
        response = await async_client.post(
            "/api/redmode/projects",
            json={"slug": slug, "display_name": slug},
            headers=headers_for("admin", "admin"),
        )
        assert response.status_code == 201
    response = await async_client.get(
        "/api/redmode/projects?limit=1&offset=1", headers=headers_for("admin", "admin")
    )
    assert response.status_code == 200
    assert response.json()["total"] == 2
    assert len(response.json()["items"]) == 1


@pytest.mark.asyncio
async def test_duplicate_project_identifier_is_rejected(async_client):
    payload = {"slug": "cliente-demo", "display_name": "Cliente Demo"}
    first = await async_client.post(
        "/api/redmode/projects", json=payload, headers=headers_for("admin", "admin")
    )
    second = await async_client.post(
        "/api/redmode/projects", json=payload, headers=headers_for("admin", "admin")
    )
    assert first.status_code == 201
    assert second.status_code == 409


@pytest.mark.asyncio
async def test_invalid_project_identifier_is_rejected(async_client):
    response = await async_client.post(
        "/api/redmode/projects",
        json={"slug": "../segredo", "display_name": "Cliente Demo"},
        headers=headers_for("admin", "admin"),
    )
    assert response.status_code == 422


@pytest.mark.asyncio
async def test_responsible_can_add_and_remove_enabled_member(async_client, fake_db):
    await grant_redmode(fake_db, "techuser")
    created = await async_client.post(
        "/api/redmode/projects",
        json={"slug": "cliente-demo", "display_name": "Cliente Demo"},
        headers=headers_for("admin", "admin"),
    )
    assert created.status_code == 201

    added = await async_client.put(
        "/api/redmode/projects/cliente-demo/members/techuser",
        headers=headers_for("admin", "admin"),
    )
    assert added.status_code == 200
    assert added.json()["members"] == ["admin", "techuser"]
    member_detail = await async_client.get(
        "/api/redmode/projects/cliente-demo", headers=headers_for("techuser")
    )
    assert member_detail.status_code == 200

    activity = await async_client.get(
        "/api/redmode/projects/cliente-demo/activity", headers=headers_for("techuser")
    )
    assert activity.status_code == 200
    assert activity.json()["items"][-1]["type"] == "member_added"

    removed = await async_client.delete(
        "/api/redmode/projects/cliente-demo/members/techuser",
        headers=headers_for("admin", "admin"),
    )
    assert removed.status_code == 200
    assert removed.json()["members"] == ["admin"]
    revoked = await async_client.get(
        "/api/redmode/projects/cliente-demo", headers=headers_for("techuser")
    )
    assert revoked.status_code == 403
    revoked_activity = await async_client.get(
        "/api/redmode/projects/cliente-demo/activity", headers=headers_for("techuser")
    )
    assert revoked_activity.status_code == 403

    activity = await async_client.get(
        "/api/redmode/projects/cliente-demo/activity", headers=headers_for("admin", "admin")
    )
    assert [item["type"] for item in activity.json()["items"]] == [
        "project_created", "member_added", "member_removed"
    ]


@pytest.mark.asyncio
async def test_only_responsible_can_change_members(async_client, fake_db):
    await grant_redmode(fake_db, "techuser")
    created = await async_client.post(
        "/api/redmode/projects",
        json={"slug": "cliente-demo", "display_name": "Cliente Demo"},
        headers=headers_for("admin", "admin"),
    )
    assert created.status_code == 201
    await async_client.put(
        "/api/redmode/projects/cliente-demo/members/techuser",
        headers=headers_for("admin", "admin"),
    )
    response = await async_client.delete(
        "/api/redmode/projects/cliente-demo/members/admin",
        headers=headers_for("techuser"),
    )
    assert response.status_code == 403


@pytest.mark.asyncio
async def test_responsible_can_change_project_phase_and_audit_it(async_client):
    created = await async_client.post(
        "/api/redmode/projects",
        json={"slug": "cliente-demo", "display_name": "Cliente Demo"},
        headers=headers_for("admin", "admin"),
    )
    assert created.status_code == 201

    changed = await async_client.put(
        "/api/redmode/projects/cliente-demo/phase",
        json={"phase": "reconnaissance"},
        headers=headers_for("admin", "admin"),
    )
    assert changed.status_code == 200
    assert changed.json()["phase"] == "reconnaissance"

    activity = await async_client.get(
        "/api/redmode/projects/cliente-demo/activity",
        headers=headers_for("admin", "admin"),
    )
    assert activity.status_code == 200
    assert activity.json()["items"][-1]["type"] == "phase_changed"
    assert activity.json()["items"][-1]["subject"] == "pre-engagement -> reconnaissance"


@pytest.mark.asyncio
async def test_project_phase_change_requires_responsible_and_valid_phase(async_client, fake_db):
    await grant_redmode(fake_db, "techuser")
    await async_client.post(
        "/api/redmode/projects",
        json={"slug": "cliente-demo", "display_name": "Cliente Demo"},
        headers=headers_for("admin", "admin"),
    )
    await async_client.put(
        "/api/redmode/projects/cliente-demo/members/techuser",
        headers=headers_for("admin", "admin"),
    )

    denied = await async_client.put(
        "/api/redmode/projects/cliente-demo/phase",
        json={"phase": "reconnaissance"},
        headers=headers_for("techuser"),
    )
    assert denied.status_code == 403

    invalid = await async_client.put(
        "/api/redmode/projects/cliente-demo/phase",
        json={"phase": "unknown"},
        headers=headers_for("admin", "admin"),
    )
    assert invalid.status_code == 422

    detail = await async_client.get(
        "/api/redmode/projects/cliente-demo",
        headers=headers_for("admin", "admin"),
    )
    assert detail.json()["phase"] == "pre-engagement"


@pytest.mark.asyncio
async def test_cannot_add_disabled_or_unenabled_user(async_client, fake_db):
    created = await async_client.post(
        "/api/redmode/projects",
        json={"slug": "cliente-demo", "display_name": "Cliente Demo"},
        headers=headers_for("admin", "admin"),
    )
    assert created.status_code == 201
    for username in ("techuser", "inactive"):
        response = await async_client.put(
            f"/api/redmode/projects/cliente-demo/members/{username}",
            headers=headers_for("admin", "admin"),
        )
        assert response.status_code == 422
    detail = await async_client.get(
        "/api/redmode/projects/cliente-demo", headers=headers_for("admin", "admin")
    )
    assert detail.json()["members"] == ["admin"]


@pytest.mark.asyncio
async def test_text_scope_publishes_version_without_review(async_client, fake_db):
    created = await async_client.post(
        "/api/redmode/projects",
        json={"slug": "cliente-demo", "display_name": "Cliente Demo"},
        headers=headers_for("admin", "admin"),
    )
    assert created.status_code == 201
    first = await async_client.post(
        "/api/redmode/projects/cliente-demo/scope/text",
        json={"text": "192.0.2.10\nTerceiros:\n203.0.113.10"},
        headers=headers_for("admin", "admin"),
    )
    assert first.status_code == 201
    scope = first.json()
    assert scope["author"] == "admin"
    assert scope["source"]["sha256"]
    assert {rule["category"] for rule in scope["rules"]} == {"client", "third_party"}
    active = await async_client.get(
        "/api/redmode/projects/cliente-demo/scope/active", headers=headers_for("admin", "admin")
    )
    assert active.status_code == 200
    assert active.json()["id"] == scope["id"]
    history = await async_client.get(
        "/api/redmode/projects/cliente-demo/scope/versions", headers=headers_for("admin", "admin")
    )
    assert history.status_code == 200
    assert [item["id"] for item in history.json()["items"]] == [scope["id"]]

    await grant_redmode(fake_db, "techuser")
    denied = await async_client.get(
        "/api/redmode/projects/cliente-demo/scope/active", headers=headers_for("techuser")
    )
    assert denied.status_code == 403


@pytest.mark.asyncio
async def test_scope_persists_offline_normalization_without_expanding_authorization(
    async_client,
):
    await async_client.post(
        "/api/redmode/projects",
        json={"slug": "cliente-demo", "display_name": "Cliente Demo"},
        headers=headers_for("admin", "admin"),
    )
    response = await async_client.post(
        "/api/redmode/projects/cliente-demo/scope/text",
        json={
            "text": (
                "192.0.2.10\n"
                "192.0.2.0/24\n"
                "HTTPS://Portal.Exämple.CO.UK:8443/login\n"
                "AS64512"
            ),
        },
        headers=headers_for("admin", "admin"),
    )
    assert response.status_code == 201
    version = response.json()
    assert version["normalization"]["mode"] == "offline"
    assert version["normalization"]["psl_version"]
    assert {rule["value"] for rule in version["rules"]} == {
        "192.0.2.10",
        "192.0.2.0/24",
        "https://portal.xn--exmple-cua.co.uk:8443/login",
    }
    ip_rule = next(rule for rule in version["rules"] if rule["kind"] == "ip")
    assert ip_rule["normalized"]["relations"][0]["canonical"] == "192.0.2.0/24"
    url_rule = next(rule for rule in version["rules"] if rule["kind"] == "url")
    assert url_rule["original_value"] == "HTTPS://Portal.Exämple.CO.UK:8443/login"
    assert url_rule["normalized"]["attributes"]["registrable_domain"] == (
        "xn--exmple-cua.co.uk"
    )
    assert [asset["value"] for asset in version["context_assets"]] == ["AS64512"]
    assert version["context_assets"][0]["executable"] is False
    assert not any(rule["kind"] == "asn" for rule in version["rules"])

    sources = await async_client.get(
        f"/api/redmode/projects/cliente-demo/scope/versions/{version['id']}/sources",
        headers=headers_for("admin", "admin"),
    )
    extraction = sources.json()["items"][0]["extraction"]
    assert extraction["rule_count"] == 3
    assert extraction["context_count"] == 1


@pytest.mark.asyncio
async def test_invalid_text_scope_keeps_previous_version(async_client):
    await async_client.post(
        "/api/redmode/projects",
        json={"slug": "cliente-demo", "display_name": "Cliente Demo"},
        headers=headers_for("admin", "admin"),
    )
    first = await async_client.post(
        "/api/redmode/projects/cliente-demo/scope/text",
        json={"text": "192.0.2.10"},
        headers=headers_for("admin", "admin"),
    )
    assert first.status_code == 201
    failed = await async_client.post(
        "/api/redmode/projects/cliente-demo/scope/text",
        json={"text": "Nenhum endereço informado."},
        headers=headers_for("admin", "admin"),
    )
    assert failed.status_code == 422
    active = await async_client.get(
        "/api/redmode/projects/cliente-demo/scope/active", headers=headers_for("admin", "admin")
    )
    assert active.json()["id"] == first.json()["id"]


@pytest.mark.asyncio
async def test_scope_history_keeps_previous_versions(async_client):
    await async_client.post(
        "/api/redmode/projects",
        json={"slug": "cliente-demo", "display_name": "Cliente Demo"},
        headers=headers_for("admin", "admin"),
    )
    versions = []
    for text in ("192.0.2.10", "192.0.2.20"):
        response = await async_client.post(
            "/api/redmode/projects/cliente-demo/scope/text",
            json={"text": text},
            headers=headers_for("admin", "admin"),
        )
        assert response.status_code == 201
        versions.append(response.json()["id"])
    history = await async_client.get(
        "/api/redmode/projects/cliente-demo/scope/versions", headers=headers_for("admin", "admin")
    )
    assert [item["id"] for item in history.json()["items"]] == versions[::-1]
    previous = await async_client.get(
        f"/api/redmode/projects/cliente-demo/scope/versions/{versions[0]}",
        headers=headers_for("admin", "admin"),
    )
    assert previous.json()["rules"][0]["value"] == "192.0.2.10"
    active = await async_client.get(
        "/api/redmode/projects/cliente-demo/scope/active", headers=headers_for("admin", "admin")
    )
    assert active.json()["id"] == versions[1]


@pytest.mark.asyncio
async def test_text_scope_limit_is_enforced(async_client):
    await async_client.post(
        "/api/redmode/projects",
        json={"slug": "cliente-demo", "display_name": "Cliente Demo"},
        headers=headers_for("admin", "admin"),
    )
    response = await async_client.post(
        "/api/redmode/projects/cliente-demo/scope/text",
        json={"text": "x" * 500001},
        headers=headers_for("admin", "admin"),
    )
    assert response.status_code == 422


@pytest.mark.asyncio
async def test_mixed_scope_upload_and_private_download(async_client, fake_db, scope_files):
    await async_client.post(
        "/api/redmode/projects",
        json={"slug": "cliente-demo", "display_name": "Cliente Demo"},
        headers=headers_for("admin", "admin"),
    )
    response = await async_client.post(
        "/api/redmode/projects/cliente-demo/scope/submit",
        data={"text": "192.0.2.10"},
        files=[
            ("files", ("scope.csv", b"category,target\nthird_party,203.0.113.10", "text/csv")),
            ("files", ("scope.json", b'{"excluded":["192.0.2.10"]}', "application/json")),
        ],
        headers=headers_for("admin", "admin"),
    )
    assert response.status_code == 201
    version = response.json()
    assert len(version["source"]["files"]) == 2
    assert {rule["value"]: rule["category"] for rule in version["rules"]} == {
        "192.0.2.10": "excluded", "203.0.113.10": "third_party"
    }
    excluded = next(rule for rule in version["rules"] if rule["value"] == "192.0.2.10")
    assert {origin["source_id"] for origin in excluded["origins"]} == {
        "text", version["source"]["files"][1]["source_id"]
    }
    file_id = version["source"]["files"][0]["id"]
    path = f"/api/redmode/projects/cliente-demo/scope/versions/{version['id']}/files/{file_id}"
    downloaded = await async_client.get(path, headers=headers_for("admin", "admin"))
    assert downloaded.status_code == 200
    assert downloaded.content == scope_files.files[file_id]
    assert downloaded.headers["cache-control"] == "private, no-store"

    await grant_redmode(fake_db, "techuser")
    denied = await async_client.get(path, headers=headers_for("techuser"))
    assert denied.status_code == 403
    denied_upload = await async_client.post(
        "/api/redmode/projects/cliente-demo/scope/submit",
        files=[("files", ("other.txt", b"198.51.100.10", "text/plain"))],
        headers=headers_for("techuser"),
    )
    assert denied_upload.status_code == 403


@pytest.mark.asyncio
async def test_scope_sources_are_materialized_and_private(
    async_client,
    fake_db,
    scope_files,
    scope_source_texts,
):
    await async_client.post(
        "/api/redmode/projects",
        json={"slug": "cliente-demo", "display_name": "Cliente Demo"},
        headers=headers_for("admin", "admin"),
    )
    response = await async_client.post(
        "/api/redmode/projects/cliente-demo/scope/submit",
        data={"text": "192.0.2.10"},
        files=[
            ("files", ("scope.txt", b"198.51.100.10", "text/plain")),
            ("files", ("scope.txt", b"notas sem alvos", "text/plain")),
        ],
        headers=headers_for("admin", "admin"),
    )
    assert response.status_code == 201
    version = response.json()
    path = f"/api/redmode/projects/cliente-demo/scope/versions/{version['id']}/sources"
    listed = await async_client.get(path, headers=headers_for("admin", "admin"))
    assert listed.status_code == 200
    sources = listed.json()["items"]
    assert [item["type"] for item in sources] == ["text", "file", "file"]
    assert [item["name"] for item in sources] == ["Texto colado", "scope.txt", "scope.txt"]
    assert len({item["source_id"] for item in sources}) == 3
    assert [item["extraction"]["rule_count"] for item in sources] == [1, 1, 0]
    assert sources[2]["extraction"]["warnings"] == ["no_targets_detected"]
    assert all(item["representation"] == {"available": True, "materialized": True} for item in sources)
    assert sources[0]["original"] == {"available": False, "file_id": None}
    assert all(item["original"]["available"] for item in sources[1:])
    assert len(scope_files.files) == 2
    assert len(scope_source_texts.files) == 3

    file_source = sources[1]
    detail = await async_client.get(
        f"{path}/{file_source['source_id']}",
        headers=headers_for("admin", "admin"),
    )
    assert detail.status_code == 200
    detail_body = detail.json()
    assert detail_body["source_id"] == file_source["source_id"]
    assert detail_body["representation"]["content"] == "198.51.100.10"
    assert detail_body["representation"]["total_characters"] == 13
    assert detail_body["rules"]["total"] == 1
    representation = await async_client.get(
        f"{path}/{file_source['source_id']}/representation",
        headers=headers_for("admin", "admin"),
    )
    assert representation.status_code == 200
    assert representation.content == b"198.51.100.10"
    assert representation.headers["cache-control"] == "private, no-store"
    assert representation.headers["x-content-type-options"] == "nosniff"

    stored = fake_db.redmode_scope_sources._data
    assert len(stored) == 3
    assert all("extracted_text" not in item for item in stored)
    await grant_redmode(fake_db, "techuser")
    denied = await async_client.get(path, headers=headers_for("techuser"))
    assert denied.status_code == 403


@pytest.mark.asyncio
async def test_scope_source_failure_rolls_back_metadata_representations_and_originals(
    async_client,
    fake_db,
    scope_files,
    scope_source_texts,
):
    await async_client.post(
        "/api/redmode/projects",
        json={"slug": "cliente-demo", "display_name": "Cliente Demo"},
        headers=headers_for("admin", "admin"),
    )
    scope_source_texts.fail_after = 1
    response = await async_client.post(
        "/api/redmode/projects/cliente-demo/scope/submit",
        data={"text": "192.0.2.10"},
        files=[("files", ("scope.txt", b"198.51.100.10", "text/plain"))],
        headers=headers_for("admin", "admin"),
    )
    assert response.status_code == 503
    assert scope_files.files == {}
    assert scope_source_texts.files == {}
    assert fake_db.redmode_scope_sources._data == []
    assert fake_db.redmode_scope_assets._data == []
    assert fake_db.redmode_scope_versions._data == []
    project = await fake_db.redmode_projects.find_one({"_id": "cliente-demo"})
    assert project.get("active_scope_version") is None
    assert project.get("scope_history", []) == []


@pytest.mark.asyncio
async def test_scope_asset_failure_rolls_back_the_entire_publication(
    async_client,
    fake_db,
    scope_files,
    scope_source_texts,
    monkeypatch,
):
    await async_client.post(
        "/api/redmode/projects",
        json={"slug": "cliente-demo", "display_name": "Cliente Demo"},
        headers=headers_for("admin", "admin"),
    )

    async def fail_asset_insert(_document):
        raise RuntimeError("asset storage unavailable")

    monkeypatch.setattr(fake_db.redmode_scope_assets, "insert_one", fail_asset_insert)
    response = await async_client.post(
        "/api/redmode/projects/cliente-demo/scope/submit",
        data={"text": "192.0.2.10"},
        files=[("files", ("scope.txt", b"198.51.100.10", "text/plain"))],
        headers=headers_for("admin", "admin"),
    )
    assert response.status_code == 503
    assert scope_files.files == {}
    assert scope_source_texts.files == {}
    assert fake_db.redmode_scope_sources._data == []
    assert fake_db.redmode_scope_assets._data == []
    assert fake_db.redmode_scope_versions._data == []


@pytest.mark.asyncio
async def test_legacy_scope_sources_are_adapted_without_rewriting_history(
    async_client,
    fake_db,
    scope_files,
):
    await async_client.post(
        "/api/redmode/projects",
        json={"slug": "cliente-demo", "display_name": "Cliente Demo"},
        headers=headers_for("admin", "admin"),
    )
    created_at = datetime.now(timezone.utc)
    version_id = "legacy-version"
    await fake_db.redmode_scope_versions.insert_one({
        "_id": version_id,
        "project_slug": "cliente-demo",
        "author": "admin",
        "created_at": created_at,
        "source": {
            "kind": "bundle",
            "text": "192.0.2.10",
            "sha256": "legacy-fingerprint",
            "files": [{
                "id": "legacy-file",
                "source_id": "legacy-source",
                "filename": "scope.txt",
                "size": 13,
                "sha256": "legacy-hash",
            }],
        },
        "rules": [{
            "value": "192.0.2.10",
            "category": "client",
            "origins": [{"source_id": "text", "line": 1}],
        }],
    })
    await fake_db.redmode_projects.update_one(
        {"_id": "cliente-demo"},
        {"$set": {"active_scope_version": version_id}, "$push": {"scope_history": version_id}},
    )
    scope_files.files["legacy-file"] = b"legacy source"
    base = f"/api/redmode/projects/cliente-demo/scope/versions/{version_id}"
    response = await async_client.get(f"{base}/sources", headers=headers_for("admin", "admin"))
    assert response.status_code == 200
    text_source, file_source = response.json()["items"]
    assert text_source["representation"] == {"available": True, "materialized": False}
    assert text_source["extraction"]["status"] == "legacy"
    assert file_source["representation"] == {"available": False, "materialized": False}
    assert "representation_not_materialized" in file_source["extraction"]["warnings"]

    text_content = await async_client.get(
        f"{base}/sources/text/representation",
        headers=headers_for("admin", "admin"),
    )
    assert text_content.status_code == 200
    assert text_content.content == b"192.0.2.10"
    absent_content = await async_client.get(
        f"{base}/sources/legacy-source/representation",
        headers=headers_for("admin", "admin"),
    )
    assert absent_content.status_code == 404
    original = await async_client.get(
        f"{base}/files/legacy-file",
        headers=headers_for("admin", "admin"),
    )
    assert original.status_code == 200
    assert original.content == b"legacy source"
    assets = await async_client.get(
        f"{base}/assets?source_id=text",
        headers=headers_for("admin", "admin"),
    )
    assert assets.status_code == 200
    assert assets.json()["total"] == 1
    assert assets.json()["items"][0]["normalized"]["canonical"] == "192.0.2.10"
    assert fake_db.redmode_scope_sources._data == []


@pytest.mark.asyncio
async def test_scope_sources_and_assets_support_server_pagination_and_filters(
    async_client,
    fake_db,
    scope_files,
):
    await async_client.post(
        "/api/redmode/projects",
        json={"slug": "cliente-demo", "display_name": "Cliente Demo"},
        headers=headers_for("admin", "admin"),
    )
    response = await async_client.post(
        "/api/redmode/projects/cliente-demo/scope/submit",
        data={
            "text": "192.0.2.10\nhttps://api.example.test/path\nAS64512",
        },
        files=[
            (
                "files",
                (
                    "scope.txt",
                    b"Excluido: 192.0.2.10\nExcluido: 198.51.100.0/24",
                    "text/plain",
                ),
            ),
            ("files", ("notes.txt", b"sem alvos executaveis", "text/plain")),
        ],
        headers=headers_for("admin", "admin"),
    )
    assert response.status_code == 201
    version = response.json()
    base = f"/api/redmode/projects/cliente-demo/scope/versions/{version['id']}"

    source_page = await async_client.get(
        f"{base}/sources?limit=1&offset=1",
        headers=headers_for("admin", "admin"),
    )
    assert source_page.status_code == 200
    assert source_page.json()["total"] == 3
    assert len(source_page.json()["items"]) == 1
    file_source_id = source_page.json()["items"][0]["source_id"]

    inventory = await async_client.get(
        f"{base}/assets?limit=2&offset=1",
        headers=headers_for("admin", "admin"),
    )
    assert inventory.status_code == 200
    assert inventory.json()["total"] == 4
    assert len(inventory.json()["items"]) == 2
    assert inventory.json()["totals"]["by_kind"] == {
        "ip": 1,
        "cidr": 1,
        "domain": 0,
        "url": 1,
        "asn": 1,
    }
    assert len(fake_db.redmode_scope_assets._data) == 4

    excluded = await async_client.get(
        f"{base}/assets?category=excluded",
        headers=headers_for("admin", "admin"),
    )
    assert excluded.json()["total"] == 2
    assert {item["kind"] for item in excluded.json()["items"]} == {"ip", "cidr"}
    merged_ip = next(item for item in excluded.json()["items"] if item["kind"] == "ip")
    assert len(merged_ip["origins"]) == 2

    asn = await async_client.get(
        f"{base}/assets?kind=asn",
        headers=headers_for("admin", "admin"),
    )
    assert asn.json()["total"] == 1
    assert asn.json()["items"][0]["executable"] is False
    assert asn.json()["items"][0]["value"] == "AS64512"

    by_source = await async_client.get(
        f"{base}/assets?source_id={file_source_id}",
        headers=headers_for("admin", "admin"),
    )
    assert by_source.json()["total"] == 2
    assert {item["kind"] for item in by_source.json()["items"]} == {"ip", "cidr"}
    searched = await async_client.get(
        f"{base}/assets?q=API.EXAMPLE",
        headers=headers_for("admin", "admin"),
    )
    assert searched.json()["total"] == 1
    literal_regex = await async_client.get(
        f"{base}/assets?q=%5B",
        headers=headers_for("admin", "admin"),
    )
    assert literal_regex.status_code == 200
    assert literal_regex.json()["total"] == 0

    text_detail = await async_client.get(
        f"{base}/sources/text?content_limit=5&rule_limit=1",
        headers=headers_for("admin", "admin"),
    )
    assert text_detail.status_code == 200
    assert text_detail.json()["representation"]["content"] == "192.0"
    assert text_detail.json()["representation"]["truncated"] is True
    assert text_detail.json()["rules"]["total"] == 2
    assert len(text_detail.json()["rules"]["items"]) == 1
    text_page = await async_client.get(
        f"{base}/sources/text/representation?offset=6&limit=8",
        headers=headers_for("admin", "admin"),
    )
    assert text_page.content == b"2.10\nhtt"
    assert text_page.headers["x-content-offset"] == "6"
    assert text_page.headers["x-content-truncated"] == "true"

    await grant_redmode(fake_db, "techuser")
    denied = await async_client.get(
        f"{base}/assets",
        headers=headers_for("techuser"),
    )
    assert denied.status_code == 403


@pytest.mark.asyncio
async def test_invalid_file_does_not_replace_active_scope(async_client, scope_files):
    await async_client.post(
        "/api/redmode/projects",
        json={"slug": "cliente-demo", "display_name": "Cliente Demo"},
        headers=headers_for("admin", "admin"),
    )
    first = await async_client.post(
        "/api/redmode/projects/cliente-demo/scope/text",
        json={"text": "192.0.2.10"},
        headers=headers_for("admin", "admin"),
    )
    failed = await async_client.post(
        "/api/redmode/projects/cliente-demo/scope/submit",
        files=[("files", ("scope.pdf", b"not a pdf", "application/pdf"))],
        headers=headers_for("admin", "admin"),
    )
    assert failed.status_code == 422
    assert scope_files.files == {}
    active = await async_client.get(
        "/api/redmode/projects/cliente-demo/scope/active", headers=headers_for("admin", "admin")
    )
    assert active.json()["id"] == first.json()["id"]


@pytest.mark.asyncio
async def test_scope_upload_limits_are_enforced(async_client, scope_files, monkeypatch):
    from config import settings
    await async_client.post(
        "/api/redmode/projects",
        json={"slug": "cliente-demo", "display_name": "Cliente Demo"},
        headers=headers_for("admin", "admin"),
    )
    monkeypatch.setattr(settings, "redmode_scope_max_file_bytes", 4)
    response = await async_client.post(
        "/api/redmode/projects/cliente-demo/scope/submit",
        files=[("files", ("scope.txt", b"192.0.2.10", "text/plain"))],
        headers=headers_for("admin", "admin"),
    )
    assert response.status_code == 413
    assert scope_files.files == {}


@pytest.mark.asyncio
async def test_file_storage_failure_does_not_publish_scope(async_client, scope_files):
    await async_client.post(
        "/api/redmode/projects",
        json={"slug": "cliente-demo", "display_name": "Cliente Demo"},
        headers=headers_for("admin", "admin"),
    )
    scope_files.fail_save = True
    response = await async_client.post(
        "/api/redmode/projects/cliente-demo/scope/submit",
        files=[("files", ("scope.txt", b"192.0.2.10", "text/plain"))],
        headers=headers_for("admin", "admin"),
    )
    assert response.status_code == 503
    active = await async_client.get(
        "/api/redmode/projects/cliente-demo/scope/active", headers=headers_for("admin", "admin")
    )
    assert active.status_code == 404


@pytest.mark.asyncio
async def test_document_scope_has_position_and_preserves_previous_on_no_text(async_client, scope_files):
    await async_client.post(
        "/api/redmode/projects",
        json={"slug": "cliente-demo", "display_name": "Cliente Demo"},
        headers=headers_for("admin", "admin"),
    )
    document = Document()
    document.add_paragraph("198.51.100.10")
    output = io.BytesIO()
    document.save(output)
    published = await async_client.post(
        "/api/redmode/projects/cliente-demo/scope/submit",
        files=[("files", ("scope.docx", output.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))],
        headers=headers_for("admin", "admin"),
    )
    assert published.status_code == 201
    assert published.json()["rules"][0]["origins"][0]["position"] == "parágrafo 1"

    blank = Document()
    blank_output = io.BytesIO()
    blank.save(blank_output)
    failed = await async_client.post(
        "/api/redmode/projects/cliente-demo/scope/submit",
        files=[("files", ("blank.docx", blank_output.getvalue(), "application/octet-stream"))],
        headers=headers_for("admin", "admin"),
    )
    assert failed.status_code == 422
    assert failed.json()["detail"] == "scope_file_no_text"
    active = await async_client.get(
        "/api/redmode/projects/cliente-demo/scope/active", headers=headers_for("admin", "admin")
    )
    assert active.json()["id"] == published.json()["id"]
